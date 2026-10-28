from __future__ import annotations

import argparse
import hashlib
import json
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any


SUPPORTED_SUFFIXES = {".docx", ".md", ".pdf", ".txt"}
MANIFEST_VERSION = 1


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def fingerprint(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def split_document(text: str) -> list[str]:
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=100,
        add_start_index=False,
    )
    return splitter.split_text(text)


def make_chunk_records(doc_id: str, chunks: list[str]) -> list[dict[str, str]]:
    occurrences: Counter[str] = Counter()
    records = []
    for chunk in chunks:
        chunk_hash = fingerprint(chunk)
        occurrence = occurrences[chunk_hash]
        occurrences[chunk_hash] += 1
        vector_id = fingerprint(f"{doc_id}\0{chunk_hash}\0{occurrence}")
        records.append(
            {
                "hash": chunk_hash,
                "vector_id": vector_id,
                "text": chunk,
            }
        )
    return records


def diff_chunk_records(
    old_records: list[dict[str, str]], new_records: list[dict[str, str]]
) -> tuple[list[dict[str, str]], list[dict[str, str]], list[str]]:
    old_ids = {record["vector_id"] for record in old_records}
    new_ids = {record["vector_id"] for record in new_records}
    added = [record for record in new_records if record["vector_id"] not in old_ids]
    reused = [record for record in new_records if record["vector_id"] in old_ids]
    deleted = [record["vector_id"] for record in old_records if record["vector_id"] not in new_ids]
    return added, reused, deleted


def load_manifest(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"version": MANIFEST_VERSION, "documents": {}}

    manifest = json.loads(path.read_text(encoding="utf-8"))
    if manifest.get("version") != MANIFEST_VERSION or not isinstance(
        manifest.get("documents"), dict
    ):
        raise ValueError(f"Unsupported or invalid manifest: {path}")
    return manifest


def save_manifest(path: Path, manifest: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_suffix(path.suffix + ".tmp")
    temporary_path.write_text(
        json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8"
    )
    temporary_path.replace(path)


def iter_documents(source_dir: Path):
    for path in sorted(source_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
            yield path


def read_document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(
            "\n".join("\t".join(cell.text for cell in row.cells) for row in table.rows)
            for table in document.tables
        )
        return "\n".join(parts)
    return path.read_text(encoding="utf-8-sig")


def save_uploaded_files(source_dir: Path, uploaded_files: list[Any]) -> list[str]:
    prepared_files = []
    seen_names: set[str] = set()
    for uploaded_file in uploaded_files:
        file_name = Path(uploaded_file.name).name
        if Path(file_name).suffix.lower() not in SUPPORTED_SUFFIXES:
            raise ValueError(f"Unsupported file type: {file_name}")
        if file_name.casefold() in seen_names:
            raise ValueError(f"Duplicate filename in upload: {file_name}")
        seen_names.add(file_name.casefold())
        prepared_files.append((file_name, uploaded_file.getvalue()))

    source_dir.mkdir(parents=True, exist_ok=True)
    for file_name, contents in prepared_files:
        (source_dir / file_name).write_bytes(contents)
    return [file_name for file_name, _ in prepared_files]


def sync_documents(source_dir: Path, manifest_path: Path, collection: Any) -> dict[str, int]:
    manifest = load_manifest(manifest_path)
    documents = manifest["documents"]
    seen_doc_ids: set[str] = set()
    stats = {
        "indexed": 0,
        "skipped": 0,
        "added_chunks": 0,
        "reused_chunks": 0,
        "deleted_chunks": 0,
        "deleted_documents": 0,
    }

    for path in iter_documents(source_dir):
        doc_id = path.relative_to(source_dir).as_posix()
        seen_doc_ids.add(doc_id)
        text = normalize_text(read_document_text(path))
        document_hash = fingerprint(text)
        previous = documents.get(doc_id)

        if previous and previous["hash"] == document_hash:
            stats["skipped"] += 1
            continue

        chunks = split_document(text)
        new_records = make_chunk_records(doc_id, chunks)
        old_records = previous["chunks"] if previous else []
        added, reused, deleted_ids = diff_chunk_records(old_records, new_records)

        if added:
            collection.upsert(
                ids=[record["vector_id"] for record in added],
                documents=[record["text"] for record in added],
                metadatas=[
                    {
                        "doc_id": doc_id,
                        "source": doc_id,
                        "chunk_hash": record["hash"],
                    }
                    for record in added
                ],
            )
        if deleted_ids:
            collection.delete(ids=deleted_ids)

        documents[doc_id] = {
            "hash": document_hash,
            "chunks": [
                {"hash": record["hash"], "vector_id": record["vector_id"]}
                for record in new_records
            ],
        }
        stats["indexed"] += 1
        stats["added_chunks"] += len(added)
        stats["reused_chunks"] += len(reused)
        stats["deleted_chunks"] += len(deleted_ids)

    for doc_id in set(documents) - seen_doc_ids:
        old_records = documents.pop(doc_id)["chunks"]
        vector_ids = [record["vector_id"] for record in old_records]
        if vector_ids:
            collection.delete(ids=vector_ids)
        stats["deleted_chunks"] += len(vector_ids)
        stats["deleted_documents"] += 1

    save_manifest(manifest_path, manifest)
    return stats


def create_collection(data_dir: Path):
    import chromadb
    from chromadb.utils.embedding_functions import SentenceTransformerEmbeddingFunction

    embedding_function = SentenceTransformerEmbeddingFunction(
        model_name="all-MiniLM-L6-v2"
    )
    client = chromadb.PersistentClient(path=str(data_dir / "chroma"))
    return client.get_or_create_collection(
        name="documents", embedding_function=embedding_function
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Incrementally index local text files.")
    parser.add_argument("--source", type=Path, default=Path("documents"))
    parser.add_argument("--data-dir", type=Path, default=Path(".rag_index"))
    args = parser.parse_args()

    source_dir = args.source.resolve()
    data_dir = args.data_dir.resolve()
    if not source_dir.is_dir():
        parser.error(f"Source directory does not exist: {source_dir}")

    data_dir.mkdir(parents=True, exist_ok=True)
    collection = create_collection(data_dir)
    stats = sync_documents(source_dir, data_dir / "manifest.json", collection)
    print(
        "Index complete: "
        f"{stats['indexed']} indexed, {stats['skipped']} unchanged, "
        f"{stats['added_chunks']} chunks added, {stats['reused_chunks']} reused, "
        f"{stats['deleted_chunks']} deleted, "
        f"{stats['deleted_documents']} documents removed."
    )


if __name__ == "__main__":
    main()