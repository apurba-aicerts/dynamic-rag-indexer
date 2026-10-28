import unittest
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from indexer import (
    SUPPORTED_SUFFIXES,
    diff_chunk_records,
    make_chunk_records,
    normalize_text,
    read_document_text,
    save_uploaded_files,
    sync_documents,
)


class FakeCollection:
    def __init__(self):
        self.records = {}

    def upsert(self, ids, documents, metadatas):
        for vector_id, document, metadata in zip(ids, documents, metadatas):
            self.records[vector_id] = {"document": document, "metadata": metadata}

    def delete(self, ids):
        for vector_id in ids:
            self.records.pop(vector_id, None)


class FakeUpload:
    def __init__(self, name, contents):
        self.name = name
        self.contents = contents

    def getvalue(self):
        return self.contents


class IndexerTests(unittest.TestCase):
    def test_normalize_text_handles_unicode_and_line_endings(self):
        self.assertEqual(normalize_text("cafe\u0301\r\ntext\r"), "caf\u00e9\ntext")

    def test_diff_reuses_unchanged_chunks_and_detects_added_and_deleted(self):
        old_records = make_chunk_records("notes.txt", ["keep", "remove"])
        new_records = make_chunk_records("notes.txt", ["keep", "add"])

        added, reused, deleted = diff_chunk_records(old_records, new_records)

        self.assertEqual([record["text"] for record in added], ["add"])
        self.assertEqual([record["text"] for record in reused], ["keep"])
        self.assertEqual(deleted, [old_records[1]["vector_id"]])

    def test_duplicate_chunks_get_distinct_stable_ids(self):
        first = make_chunk_records("notes.txt", ["repeat", "repeat"])
        second = make_chunk_records("notes.txt", ["repeat", "repeat"])

        self.assertNotEqual(first[0]["vector_id"], first[1]["vector_id"])
        self.assertEqual(
            [record["vector_id"] for record in first],
            [record["vector_id"] for record in second],
        )

    def test_upload_saves_supported_files_using_only_the_filename(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            source_dir = Path(temporary_directory) / "documents"

            saved = save_uploaded_files(
                source_dir, [FakeUpload("folder\\notes.TXT", b"hello")]
            )

            self.assertEqual(saved, ["notes.TXT"])
            self.assertEqual((source_dir / "notes.TXT").read_bytes(), b"hello")
            self.assertTrue({".pdf", ".docx"}.issubset(SUPPORTED_SUFFIXES))

    def test_upload_rejects_unsupported_and_duplicate_filenames(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            source_dir = Path(temporary_directory) / "documents"

            with self.assertRaisesRegex(ValueError, "Unsupported file type"):
                save_uploaded_files(source_dir, [FakeUpload("image.png", b"x")])
            with self.assertRaisesRegex(ValueError, "Duplicate filename"):
                save_uploaded_files(
                    source_dir,
                    [FakeUpload("Notes.txt", b"one"), FakeUpload("notes.TXT", b"two")],
                )

    def test_pdf_extraction_combines_page_text(self):
        pages = [
            SimpleNamespace(extract_text=lambda: "first page"),
            SimpleNamespace(extract_text=lambda: "second page"),
        ]
        with patch("pypdf.PdfReader", return_value=SimpleNamespace(pages=pages)):
            self.assertEqual(
                read_document_text(Path("sample.pdf")), "first page\nsecond page"
            )

    def test_docx_extraction_includes_paragraphs_and_tables(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "sample.docx"
            document = Document()
            document.add_paragraph("A paragraph")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "First cell"
            table.cell(0, 1).text = "Second cell"
            document.save(path)

            extracted = read_document_text(path)

            self.assertIn("A paragraph", extracted)
            self.assertIn("First cell\tSecond cell", extracted)

    def test_sync_updates_only_changed_chunks_and_removes_deleted_documents(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            source_dir = root / "documents"
            source_dir.mkdir()
            source_file = source_dir / "notes.txt"
            manifest_path = root / "manifest.json"
            collection = FakeCollection()

            with patch("indexer.split_document", side_effect=lambda text: text.split("|")):
                source_file.write_text("keep|remove", encoding="utf-8")
                first = sync_documents(source_dir, manifest_path, collection)
                self.assertEqual(first["added_chunks"], 2)

                unchanged = sync_documents(source_dir, manifest_path, collection)
                self.assertEqual(unchanged["skipped"], 1)

                source_file.write_text("keep|add", encoding="utf-8")
                changed = sync_documents(source_dir, manifest_path, collection)
                self.assertEqual(changed["added_chunks"], 1)
                self.assertEqual(changed["reused_chunks"], 1)
                self.assertEqual(changed["deleted_chunks"], 1)
                self.assertEqual(
                    {record["document"] for record in collection.records.values()},
                    {"keep", "add"},
                )

                source_file.unlink()
                deleted = sync_documents(source_dir, manifest_path, collection)
                self.assertEqual(deleted["deleted_documents"], 1)
                self.assertEqual(collection.records, {})


if __name__ == "__main__":
    unittest.main()